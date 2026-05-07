# update_docx.py — tekazashi_guide.docx を更新するスクリプト
# 変更内容:
#   1. 各タブの機能テーブル「統計」行の説明文を更新
#   2. よく使う操作に 7〜9 を追加（サブタイトル・集計開始月・統計グラフ）
#   3. ★ キャッシュ注意ボックスを「データの保存について」セクション末尾に追加

import sys, io, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOCX_PATH = r'C:\Users\bench\OneDrive\Desktop\yudai\その他\手かざしアプリ\master\manuals\tekazashi_guide.docx'

doc = Document(DOCX_PATH)

# ============================================================
# 1. 統計タブの説明文を更新（Table4 の Row4）
# ============================================================
stats_new_text = (
    '人/月/年のサブタブで施光・受光の集計を表示。\n'
    '「月」タブでは棒グラフ（施光/受光/積み上げ切替）と累積折れ線グラフも確認できる。\n'
    '各人の「\U0001f4cb 履歴」ボタンで過去の施光履歴を参照。\n'
    '⚙️設定の「集計開始月」で集計対象期間を絞り込める。'
)

t4 = doc.tables[4]
stats_cell = t4.rows[4].cells[1]
# 既存テキストをクリアして新しいテキストを設定
for para in stats_cell.paragraphs:
    for run in para.runs:
        run.text = ''
if stats_cell.paragraphs:
    # 1段落目に書き込む
    p = stats_cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = stats_new_text
    else:
        p.add_run(stats_new_text)
print('✓ 統計タブの説明を更新しました')

# ============================================================
# 2. よく使う操作に 7〜9 を追加
#    paragraph[28] の直後に挿入する
# ============================================================

def make_bold_run(text, size_halfpt=22):
    """太字ランのXML要素を返す"""
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    b = OxmlElement('w:b')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(size_halfpt))
    szcs = OxmlElement('w:szCs')
    szcs.set(qn('w:val'), str(size_halfpt))
    rpr.append(b); rpr.append(sz); rpr.append(szcs)
    t = OxmlElement('w:t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(rpr); r.append(t)
    return r

def make_normal_run(text, size_halfpt=20):
    """通常ランのXML要素を返す"""
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(size_halfpt))
    szcs = OxmlElement('w:szCs')
    szcs.set(qn('w:val'), str(size_halfpt))
    rpr.append(sz); rpr.append(szcs)
    t = OxmlElement('w:t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(rpr); r.append(t)
    return r

def make_title_para(number_str, title_str):
    """番号＋タイトルの太字パラグラフ要素を返す"""
    p = OxmlElement('w:p')
    p.append(make_bold_run(number_str, 22))
    p.append(make_bold_run(title_str,  22))
    return p

def make_body_para(text):
    """説明文の通常パラグラフ要素を返す"""
    p = OxmlElement('w:p')
    p.append(make_normal_run(text, 20))
    return p

# 追加する項目
new_ops = [
    ('7　', 'サブタイトルのカスタマイズ',
     '右上の⚙️設定ボタン → 「サブタイトル」欄に文字を入力して「保存」をタップします。'
     'ヘッダーに表示される一言を自分用にカスタマイズできます。'),
    ('8　', '集計開始月の設定',
     '右上の⚙️設定ボタン → 「集計開始月」で月を選んで「保存」をタップします。'
     '統計タブや施光・受光タブの累計が、設定した月以降のデータで集計されます。'
     '「全期間に戻す」で解除できます。'),
    ('9　', '統計グラフの切り替え',
     '統計タブ → 「月」サブタブを開くと月別の棒グラフと累積折れ線グラフが表示されます。'
     '棒グラフ上部の「施光」「受光」「積み上げ」ボタンで表示を切り替えられます。'),
]

# paragraph[28] の直後に挿入（逆順で addnext）
anchor = doc.paragraphs[28]._element
for number_str, title_str, body_str in reversed(new_ops):
    anchor.addnext(make_body_para(body_str))
    anchor.addnext(make_title_para(number_str, title_str))

print('✓ よく使う操作に 7〜9 を追加しました')

# ============================================================
# 3. ★ キャッシュ注意ボックスを Table7 の直後に追加
# ============================================================
# Table7 の構造をコピーして新しい内容で作成する
cache_text = (
    '★ アップデート後に画面が古いままの場合：\n'
    'メンテナンス（機能追加）後しばらくは、ブラウザがキャッシュ（保存された古いファイル）を\n'
    '表示することがあります。正しく表示されない場合は以下をお試しください。\n'
    '\n'
    '・iPhone / Safari：「設定」→「Safari」→「履歴とWebサイトのデータを消去」\n'
    '・Android / Chrome：右上の⋮メニュー → 「設定」→「プライバシー」→「閲覧データを削除」\n'
    '\n'
    'またはプライベートブラウズモードで開くと、常に最新バージョンが表示されます。'
)

# Table7 の XML 要素を取得してコピーする
t7 = doc.tables[7]
t7_elem = t7._element
new_table = copy.deepcopy(t7_elem)

# コピーしたテーブルのテキストを差し替える
# （セル内の最初のパラグラフランにまとめて書き込む）
ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
tc_list = new_table.findall(f'.//{{{ns}}}tc')
if tc_list:
    tc = tc_list[0]
    t_elems = tc.findall(f'.//{{{ns}}}t')
    # 最初の <w:t> にまとめて書き込み、残りを空にする
    if t_elems:
        t_elems[0].text = cache_text
        for te in t_elems[1:]:
            te.text = ''

# Table7 の直後に挿入
t7_elem.addnext(new_table)
print('✓ ★ キャッシュ注意ボックスを追加しました')

# ============================================================
# 保存
# ============================================================
doc.save(DOCX_PATH)
print('\n完了：tekazashi_guide.docx を更新しました')
